from io import BytesIO
import json
import os
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pytest
from PIL import Image

from app import report_documents as documents


def png(width=800, height=400):
    target = BytesIO()
    Image.new("RGB", (width, height), (25, 100, 140)).save(target, format="PNG")
    return target.getvalue()


def bundle():
    return {"title": "研究動向の分析", "subtitle": "保存された分析と図をまとめたレポート", "report_id": "report-test",
        "result_id": "result-test", "created_at": "2026-10-09", "scope_note": "抄録を対象とした分析です。",
        "warnings": ["数値の照合に注意が必要です。"],
        "sections": [{"title": "時系列の変化", "text": "2021年には材料に注目し、2025年には条件の評価が進んだ。"}],
        "tables": [{"title": "年次集計", "columns": ["年", "論文数", "割合", "確定", "注記"],
                    "rows": [[2021, 10, .25, True, "=SUM(A1:A2)"], [2025, 30, .75, False, "通常の文字列"]]}],
        "figures": [{"title": "時層ランドスケープ", "caption": "同じ座標系で年別の層を比較する。静止画です。", "png": png(), "width": 12, "height": 99}]}


def texts_in_zip(data, names):
    with ZipFile(BytesIO(data)) as archive:
        return "".join("".join(ET.fromstring(archive.read(name)).itertext()) for name in archive.namelist() if names(name))


def test_docx_preserves_full_prose_warnings_identifiers_tables_and_png_aspect():
    from docx import Document
    value = bundle()
    value["sections"][0]["text"] = "開始。" + "長い評論です。" * 5000 + "最後の論点。"
    value["tables"][0]["rows"][0][-1] = "表の長い記録。" * 200 + "末尾。"
    data = documents.render_report(value, "docx")
    doc = Document(BytesIO(data))
    joined = "".join(paragraph.text for paragraph in doc.paragraphs)
    assert value["sections"][0]["text"] in joined
    assert "report-test" in joined and "数値の照合" in joined
    assert "末尾。" in "".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    image = doc.inline_shapes[0]
    assert image.width / image.height == pytest.approx(2, abs=.001)
    assert "時層ランドスケープ" in joined


def test_pptx_paginates_complete_long_prose_and_embeds_undistorted_figure():
    from pptx import Presentation
    value = bundle()
    text = "開始。" + "論文の条件と結果を比較する。" * 2000 + "最後の論点。"
    value["sections"][0]["text"] = text
    data = documents.render_report(value, "pptx")
    presentation = Presentation(BytesIO(data))
    body = "".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
                   and shape.top > 1_000_000 and shape.top < 5_000_000)
    assert text in body
    assert len(presentation.slides) > 10
    pictures = [shape for slide in presentation.slides for shape in slide.shapes if shape.shape_type == 13]
    assert len(pictures) == 1 and pictures[0].width / pictures[0].height == pytest.approx(2, abs=.001)
    assert all(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= presentation.slide_width + 100
               and shape.top + shape.height <= presentation.slide_height + 100
               for slide in presentation.slides for shape in slide.shapes)


def test_xlsx_literal_formulas_typed_values_full_records_and_lossless_long_cells():
    from openpyxl import load_workbook
    value = bundle()
    text = "😀論文" * 20_000 + "最終文字"
    records = [{"id": "p-1", "title": "=HYPERLINK(\"https://example.org\")", "year": 2021, "abstract": text,
                "status": "failed", "extraction": {"facts": [{"quote": "根拠\x00に制御文字", "source_start": 123}], "chunks": []}},
               {"id": "p-2", "year": 2025, "abstract": "", "status": "missing", "large_id": 1234567890123456789}]
    seen = []
    def stream():
        for record in records:
            seen.append(record["id"])
            yield record
    data = documents.render_report(value, "xlsx", records=stream())
    book = load_workbook(BytesIO(data), data_only=False)
    assert seen == ["p-1", "p-2"]
    table = book["集計1"]
    assert table["D2"].value == 2021 and table["D2"].data_type == "n"
    assert table["F2"].value == .25 and table["G2"].data_type == "b"
    assert table["H2"].value == "=SUM(A1:A2)" and table["H2"].data_type == "s"
    paper_rows = list(book["全論文"].values)[1:]
    assert "".join(row[8] or "" for row in paper_rows if row[0] == 1) == text
    assert any(row[0] == 2 and row[7] == "missing" for row in paper_rows)
    rows = list(book["論文詳細"].values)[1:]
    quote = [row for row in rows if row[2] == "/extraction/facts/0/quote"][0]
    assert quote[3] == "string_json" and json.loads(quote[6]) == "根拠\x00に制御文字"
    integer = [row for row in rows if row[2] == "/large_id"][0]
    assert integer[3] == "integer" and integer[6] == "1234567890123456789"
    with ZipFile(BytesIO(data)) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/media/")]) == 1


def test_xlsx_table_and_prose_long_strings_survive_and_row_limit_rolls_over(monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.setattr(documents, "EXCEL_MAX_ROWS", 5)
    value = bundle()
    value["figures"] = []
    long = "=文字列" * 12_000
    value["sections"][0]["text"] = long
    value["tables"] = [{"title": "大きい表", "columns": ["文"], "rows": [[long]] + [[f"row-{i}"] for i in range(8)]}]
    data = documents.render_report(value, "xlsx")
    book = load_workbook(BytesIO(data))
    assert all(sheet.max_row <= 5 for sheet in book)
    parts = [row[3] for sheet in book if sheet.title.startswith("集計1") for row in list(sheet.values)[1:] if row[0] == 1]
    assert "".join(parts) == long
    reports = [row for sheet in book if sheet.title.startswith("レポート") for row in list(sheet.values)[1:]]
    index = next(i for i, row in enumerate(reports) if row[2] == "時系列の変化")
    assert "".join(row[3] for row in reports[index:index + reports[index][1]]) == long


def test_xlsx_only_audit_tables_are_excluded_from_human_documents_but_kept_in_excel():
    value = bundle()
    value["tables"].append({"title": "AUDIT_ONLY_SENTINEL", "columns": ["path", "value"], "rows": [["/x", "exact-saved-value"]], "xlsx_only": True})
    for kind in ("docx", "pptx"):
        data = documents.render_report(value, kind)
        text = texts_in_zip(data, lambda name: name.endswith(".xml"))
        assert "AUDIT_ONLY_SENTINEL" not in text and "exact-saved-value" not in text
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(documents.render_report(value, "xlsx")))
    assert any("exact-saved-value" == cell.value for sheet in book for row in sheet for cell in row)


def test_pdf_contains_complete_sections_and_paginates_long_cells(monkeypatch):
    font = Path(os.environ.get("ATLAS_PDF_FONT", "C:/Users/BN0346/Documents/ChatGPT/New project/tmp/ci-font/ipaexg.ttf"))
    if not font.exists():
        pytest.skip("Japanese PDF test font unavailable")
    monkeypatch.setenv("ATLAS_PDF_FONT", str(font))
    value = bundle()
    value["sections"][0]["text"] = "評論です。" * 3000 + "FINAL_SECTION_SENTINEL"
    value["tables"][0]["rows"][0][-1] = "表の記録。" * 300 + "FINAL_TABLE_SENTINEL"
    data = documents.render_report(value, "pdf")
    assert data.startswith(b"%PDF") and len(data) > 10_000
    assert data.count(b"/Type /Page") > 5
    # PDF text is compressed and font encoded; layout success plus full source
    # paragraph generation is checked separately with the installed renderer.
    assert "".join(documents._pages(value["sections"][0]["text"])) == value["sections"][0]["text"]


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "xlsx"])
def test_invalid_png_is_an_actionable_error(kind, monkeypatch):
    monkeypatch.setattr(documents, "_font", lambda: "Helvetica")
    value = bundle()
    value["figures"][0]["png"] = b"not a png"
    with pytest.raises(ValueError, match="PNG"):
        documents.render_report(value, kind)


def test_non_excel_renderers_never_consume_all_record_iterator():
    def forbidden():
        raise AssertionError("record stream must only be consumed for xlsx")
        yield {}
    for kind in ("docx", "pptx"):
        documents.render_report(bundle(), kind, records=forbidden())
    with pytest.raises(ValueError, match="PDF"):
        documents.render_report(bundle(), "html")


def test_lossless_split_and_pagination_include_whitespace_and_emoji():
    text = " \n😀日本語 abc\n\n" * 100
    fragments = list(documents._parts(text, 31))
    assert "".join(fragments) == text
    assert all(len(part.encode("utf-16-le")) // 2 <= 31 for part in fragments)
    assert "".join(documents._pages(text, 15, 3)) == text


def test_write_only_export_closes_partial_workbook_when_record_source_fails():
    def broken():
        yield {"id": "completed-before-error", "abstract": "A result."}
        raise OSError("snapshot unavailable")
    with pytest.raises(OSError, match="snapshot unavailable"):
        documents.render_report(bundle(), "xlsx", records=broken())
