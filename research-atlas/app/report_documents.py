"""Portable document exports from saved report content and rendered figures.

No renderer calls an LLM. The bundle is a presentation snapshot; XLSX may also
receive a streaming iterator containing every corpus record, including failures.
"""
from __future__ import annotations

from html import escape
from io import BytesIO
import json
import math
import re
import unicodedata

from .foresight_exports import _font

FORMATS = {"pdf", "docx", "xlsx", "pptx"}
EXCEL_MAX_ROWS = 1_048_576
EXCEL_CELL_UNITS = 30_000
OFFICE_FONT = "Yu Gothic"
_INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return str(value)


def _xml(value):
    # XML cannot contain these characters. Show an explicit escape, rather than
    # silently discard it; the XLSX audit retains a reversible JSON form.
    return _INVALID_XML.sub(lambda match: "\\u%04x" % ord(match.group()), _text(value))


def _parts(text, limit=EXCEL_CELL_UNITS):
    """Lossless splits under Excel's UTF-16 length ceiling, including emoji."""
    text = _text(text)
    start = units = 0
    for index, char in enumerate(text):
        size = 2 if ord(char) > 0xFFFF else 1
        if units + size > limit:
            yield text[start:index]
            start, units = index, 0
        units += size
    yield text[start:]


def _pages(text, columns=88, lines=16):
    """Approximate CJK line width conservatively; never strip source text."""
    text = _xml(text)
    start = used = line = 0
    for index, char in enumerate(text):
        width = 2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
        if char == "\n":
            line, used = line + 1, 0
        elif used + width > columns:
            line, used = line + 1, width
        else:
            used += width
        if line >= lines:
            yield text[start:index + 1]
            start, used, line = index + 1, 0, 0
    if start < len(text) or not text:
        yield text[start:]


def _meta(bundle):
    return f"作成日時: {_text(bundle.get('created_at'))}\nReport: {_text(bundle.get('report_id'))}\nAnalysis: {_text(bundle.get('result_id'))}"


def _blocks(bundle):
    yield "レポート情報", _meta(bundle)
    if bundle.get("subtitle"):
        yield "概要", _text(bundle["subtitle"])
    if bundle.get("scope_note"):
        yield "対象範囲・読み方", _text(bundle["scope_note"])
    if bundle.get("warnings"):
        yield "注意事項", "\n".join(_text(item) for item in bundle["warnings"])
    for section in bundle.get("sections", []):
        yield _text(section.get("title")), _text(section.get("text"))


def _panels(table, *, max_columns=5, cell_chars=180):
    columns = [_text(value) for value in table.get("columns", [])]
    rows = table.get("rows", [])
    if not columns:
        return
    for offset in range(0, len(columns), max_columns):
        labels = columns[offset:offset + max_columns]
        panel_rows = []
        for row_index, row in enumerate(rows, 1):
            pieces = [list(_parts(_text(row[index]) if index < len(row) else "", cell_chars))
                      for index in range(offset, offset + len(labels))]
            for part in range(max(map(len, pieces), default=1)):
                panel_rows.append([str(row_index) + (f" / 続き{part + 1}" if part else ""),
                                   *[value[part] if part < len(value) else "" for value in pieces]])
        yield ["行", *labels], panel_rows, offset


def _figure(figure):
    value = figure.get("png")
    if not isinstance(value, bytes) or not value.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError("図の画像はPNG形式のバイト列で指定してください。")
    # Actual dimensions are authoritative; do not distort a figure because a
    # browser's reported viewport and exported PNG dimensions differ.
    from PIL import Image
    try:
        with Image.open(BytesIO(value)) as image:
            width, height = image.size
            image.verify()
    except Exception:
        raise ValueError("図のPNG画像を読み込めませんでした。") from None
    if width < 1 or height < 1:
        raise ValueError("図のPNG画像サイズが不正です。")
    return value, width, height


def render_report(bundle, format, records=None):
    """Render PDF/DOCX/XLSX/PPTX; records are consumed only for XLSX."""
    kind = str(format).lower().lstrip(".")
    if kind not in FORMATS:
        raise ValueError("対応する出力形式は PDF、DOCX、XLSX、PPTX です。")
    if not isinstance(bundle, dict):
        raise ValueError("レポートの内容が不正です。")
    for figure in bundle.get("figures", []):
        _figure(figure)
    try:
        return {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "pptx": _pptx}[kind](bundle, records)
    except ImportError:
        package = {"pdf": "reportlab", "docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx"}[kind]
        raise ValueError(f"{kind.upper()}出力には {package} が必要です。requirements.txt をインストールしてください。") from None


def _pdf(bundle, _records):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
    font = _font()
    width = A4[0] - 84
    body = ParagraphStyle("body", fontName=font, fontSize=9, leading=15, wordWrap="CJK", spaceAfter=8,
                          textColor=colors.HexColor("#183147"))
    heading = ParagraphStyle("heading", parent=body, fontSize=13, leading=19, spaceBefore=14, keepWithNext=True,
                             textColor=colors.HexColor("#087F86"))
    title = ParagraphStyle("title", parent=body, fontSize=21, leading=29, spaceAfter=15)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=11)
    def p(value, style=body):
        return Paragraph(escape(_xml(value)).replace("\n", "<br/>"), style)
    story = [p("RESEARCH ATLAS", small), p(bundle.get("title", "研究レポート"), title)]
    for name, text in _blocks(bundle):
        story.append(p(name, heading))
        story.extend(p(part) for part in _pages(text, columns=100, lines=26))
    for table in bundle.get("tables", []):
        if table.get("xlsx_only"):
            continue
        story.append(p(table.get("title", "集計表"), heading))
        for headers, rows, offset in _panels(table):
            if offset:
                story.append(p(f"列 {offset + 1} 以降（行番号は原表と共通）", small))
            # Column headings can themselves be lengthy: keep the complete
            # heading outside the repeating compact table header.
            for index, header in enumerate(headers[1:], offset + 1):
                if len(header) > 60:
                    story.extend(p(part, small) for part in _pages(f"列 {index}: {header}", lines=20))
            display_headers = [header if len(header) <= 60 else f"列 {offset + index}" for index, header in enumerate(headers)]
            data = [[p(value, small) for value in display_headers], *[[p(value, small) for value in row] for row in rows]]
            first = 46
            grid = Table(data, colWidths=[first] + [(width - first) / (len(headers) - 1)] * (len(headers) - 1), repeatRows=1)
            grid.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEFEB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.extend([grid, Spacer(1, 10)])
    for figure in bundle.get("figures", []):
        raw, w, h = _figure(figure)
        scale = min(width / w, 480 / h)
        story.extend([PageBreak(), p(figure.get("title", "分析図"), heading), Image(BytesIO(raw), width=w * scale, height=h * scale)])
        story.extend(p(part, small) for part in _pages(figure.get("caption", ""), lines=24))
    target = BytesIO()
    def footer(canvas, document):
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#56697B"))
        canvas.drawString(42, 25, "Research Atlas | 保存済み分析のレポート")
        canvas.drawRightString(A4[0] - 42, 25, str(document.page))
    SimpleDocTemplate(target, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=40, bottomMargin=42,
        title=_xml(bundle.get("title", "Research Atlas")), author="Research Atlas").build(story, onFirstPage=footer, onLaterPages=footer)
    return target.getvalue()


def _docx(bundle, _records):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = OFFICE_FONT
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), OFFICE_FONT)
        style.font.color.rgb = RGBColor.from_string("183147")
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    section.header.paragraphs[0].text = "RESEARCH ATLAS | SAVED REPORT"
    footer = section.footer.paragraphs[0]
    footer.text = "Research Atlas  |  "
    number = OxmlElement("w:fldSimple")
    number.set(qn("w:instr"), "PAGE")
    footer._p.append(number)
    document.add_heading(_xml(bundle.get("title", "研究レポート")), 0)
    for name, text in _blocks(bundle):
        document.add_heading(_xml(name), 1)
        for part in _pages(text, columns=100, lines=28):
            document.add_paragraph(part)
    for source in bundle.get("tables", []):
        if source.get("xlsx_only"):
            continue
        document.add_heading(_xml(source.get("title", "集計表")), 1)
        for headers, rows, offset in _panels(source, cell_chars=800):
            if offset:
                document.add_paragraph(f"列 {offset + 1} 以降（行番号は原表と共通）")
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = _xml(text)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                for cell, text in zip(table.add_row().cells, values):
                    cell.text = _xml(text)
            document.add_paragraph()
    for figure in bundle.get("figures", []):
        raw, w, h = _figure(figure)
        scale = min(6.65 / w, 7.2 / h)
        document.add_page_break()
        document.add_heading(_xml(figure.get("title", "分析図")), 1)
        document.add_picture(BytesIO(raw), width=Inches(w * scale), height=Inches(h * scale))
        document.add_paragraph(_xml(figure.get("caption", "")))
    document.core_properties.title = _xml(bundle.get("title", "Research Atlas"))[:255]
    document.core_properties.subject = _xml(_meta(bundle))[:255]
    document.core_properties.author = "Research Atlas"
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def _xlsx(bundle, records):
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    try:
        return _xlsx_build(bundle, records, workbook)
    finally:
        # Write-only worksheets hold XML generators and temporary files. Close
        # them even if a record iterator or image export fails mid-document.
        for sheet in workbook.worksheets:
            try:
                if not sheet.closed:
                    sheet.close()
            except Exception:
                pass
            try:
                if sheet._writer:
                    sheet._writer.cleanup()
            except (OSError, AttributeError):
                pass
        workbook.close()


def _xlsx_build(bundle, records, workbook):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.drawing.image import Image
    workbook.properties.title = _xml(bundle.get("title", "Research Atlas"))
    workbook.properties.creator = "Research Atlas"
    streams = []
    used = set()
    normal_font = Font(name=OFFICE_FONT, size=10, color="183147")
    heading_font = Font(name=OFFICE_FONT, size=10, bold=True, color="FFFFFF")
    heading_fill = PatternFill("solid", fgColor="087F86")
    wrapped = Alignment(vertical="top", wrap_text=True)
    def excel_parts(value):
        if _INVALID_XML.search(value):
            value = json.dumps(value, ensure_ascii=True)
        return list(_parts(value))
    def unique(name):
        clean = re.sub(r"[\\/*?:\[\]]", "_", _xml(name)).strip("'")[:26] or "Report"
        number, result = 1, clean
        while result.casefold() in used:
            number += 1
            result = f"{clean[:25]}_{number}"
        used.add(result.casefold())
        return result
    def cell(sheet, value, header=False):
        if type(value) is float and not math.isfinite(value):
            value = str(value)
        if type(value) is int and abs(value) >= 10**15:
            value = str(value)  # Excel numbers have at most 15 significant digits.
        if isinstance(value, str):
            value = _xml(value)
        result = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            result.data_type = "s"  # Literal text, even if it starts with '='.
        result.font = heading_font if header else normal_font
        result.alignment = wrapped
        if header:
            result.fill = heading_fill
        return result
    class Rows:
        def __init__(self, name, headers):
            self.name, self.headers = name, headers
            self.sheet, self.count = None, 0
            self.open()
        def open(self):
            self.sheet = workbook.create_sheet(unique(self.name))
            self.sheet.freeze_panes = "A2"
            self.sheet.sheet_view.showGridLines = False
            for index in range(len(self.headers)):
                from openpyxl.utils import get_column_letter
                self.sheet.column_dimensions[get_column_letter(index + 1)].width = 24 if index < len(self.headers) - 1 else 70
            self.sheet.append([cell(self.sheet, value, True) for value in self.headers])
            self.count = 1
        def append(self, values):
            if self.count >= EXCEL_MAX_ROWS:
                self.open()
            self.sheet.append([cell(self.sheet, value) for value in values])
            self.count += 1
    def split_row(writer, values):
        parts = [excel_parts(value) if isinstance(value, str) else [value] for value in values]
        length = max(map(len, parts), default=1)
        for index in range(length):
            writer.append([index + 1, length, *[part[index] if index < len(part) else None for part in parts]])
    report = Rows("レポート", ["分割番号", "分割数", "項目", "内容"])
    split_row(report, ["タイトル", bundle.get("title", "研究レポート")])
    split_row(report, ["Excelの読み方", "長文は分割番号の順に連結してください。先頭が = の文字列も文字として保存します。制御文字を含むセルはJSON文字列で保持します。15桁を超える整数は文字列として保持します。"])
    for name, text in _blocks(bundle):
        split_row(report, [name, text])
    for index, table in enumerate(bundle.get("tables", []), 1):
        # Wide tables also stay within Excel's fixed 16,384-column limit.
        columns = table.get("columns", [])
        for offset in range(0, len(columns), 16_380):
            headers = [_text(value) for value in columns[offset:offset + 16_380]]
            writer = Rows(f"集計{index}", ["行番号", "分割番号", "分割数", *[header[:1000] for header in headers]])
            for row_index, row in enumerate(table.get("rows", []), 1):
                values = list(row[offset:offset + len(headers)])
                values += [None] * (len(headers) - len(values))
                pieces = [excel_parts(value) if isinstance(value, str) else [value] for value in values]
                count = max(map(len, pieces), default=1)
                for part in range(count):
                    writer.append([row_index, part + 1, count, *[value[part] if part < len(value) else None for value in pieces]])
            split_row(report, [f"集計{index}の表題", table.get("title", "")])
            for column_index, name in enumerate(headers, offset + 1):
                split_row(report, [f"集計{index}の列{column_index}", name])
    for index, figure in enumerate(bundle.get("figures", []), 1):
        raw, w, h = _figure(figure)
        writer = Rows(f"図{index}", ["分割番号", "分割数", "項目", "内容"])
        split_row(writer, ["図のタイトル", figure.get("title", "")])
        split_row(writer, ["図の説明", figure.get("caption", "")])
        split_row(writer, ["画像サイズ", f"{w} × {h} px / 元PNGを埋め込み"])
        buffer = BytesIO(raw)
        streams.append(buffer)
        image = Image(buffer)
        scale = min(1, 1120 / w, 700 / h)
        image.width, image.height = w * scale, h * scale
        writer.sheet.add_image(image, f"A{writer.count + 2}")
    if records is not None:
        papers = Rows("全論文", ["通番", "分割番号", "分割数", "論文ID", "タイトル", "年", "分野", "状態", "抄録"])
        detail = Rows("論文詳細", ["通番", "論文ID", "JSON Pointer", "型", "分割番号", "分割数", "値"])
        def visit(value, path=""):
            if isinstance(value, dict) and value:
                for key, child in value.items():
                    yield from visit(child, path + "/" + str(key).replace("~", "~0").replace("/", "~1"))
            elif isinstance(value, list) and value:
                for index, child in enumerate(value):
                    yield from visit(child, path + "/" + str(index))
            else:
                kind = ("null" if value is None else "boolean" if type(value) is bool else "integer" if type(value) is int
                        else "number" if type(value) is float else "object" if isinstance(value, dict) else "array" if isinstance(value, list) else "string")
                if isinstance(value, (dict, list)) or isinstance(value, str) and _INVALID_XML.search(value):
                    value = json.dumps(value, ensure_ascii=True, allow_nan=False)
                    kind += "_json"
                yield path, kind, value
        for ordinal, record in enumerate(records, 1):
            pid = _text(record.get("paper_id", record.get("id", "")))
            values = [pid, record.get("title", ""), record.get("year"), record.get("topic_id", ""), record.get("status", ""), record.get("abstract", "")]
            fragments = [excel_parts(value) if isinstance(value, str) else [value] for value in values]
            count = max(map(len, fragments))
            for part in range(count):
                papers.append([ordinal, part + 1, count, *[value[part] if part < len(value) else None for value in fragments]])
            for path, kind, value in visit(record):
                fragments = list(_parts(value)) if isinstance(value, str) else [value]
                for part, fragment in enumerate(fragments, 1):
                    detail.append([ordinal, pid, path, kind, part, len(fragments), fragment])
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def _pptx(bundle, _records):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.title = _xml(bundle.get("title", "Research Atlas"))[:255]
    presentation.core_properties.author = "Research Atlas"
    def text_box(slide, text, left, top, width, height, *, size=18, color="183147", bold=False):
        frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.03)
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(_xml(text).split("\n")):
            paragraph = frame.paragraphs[0] if not index else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = OFFICE_FONT
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(size * 1.2)
        return frame
    def slide(title):
        item = presentation.slides.add_slide(presentation.slide_layouts[6])
        item.background.fill.solid()
        item.background.fill.fore_color.rgb = RGBColor.from_string("F6F9FB")
        display = _xml(title)
        text_box(item, display if len(display) <= 68 else display[:65] + "…", .55, .45, 12.2, .9, size=25, color="087F86", bold=True)
        text_box(item, f"Research Atlas | 保存済み分析 | {len(presentation.slides)}", .55, 7.08, 12.2, .2, size=9, color="56697B")
        return item
    def prose(title, text):
        if len(_xml(title)) > 68:
            text = _xml(title) + "\n\n" + _xml(text)
        pages = list(_pages(text))
        for index, page in enumerate(pages):
            item = slide(title + (f"（続き {index + 1}/{len(pages)}）" if len(pages) > 1 else ""))
            text_box(item, page, .65, 1.45, 12, 5.2, size=18)
    prose("RESEARCH ATLAS", _text(bundle.get("title", "研究レポート")) + "\n\n" + _meta(bundle))
    for name, text in _blocks(bundle):
        if name != "レポート情報":
            prose(name, text)
    for source in bundle.get("tables", []):
        if source.get("xlsx_only"):
            continue
        name = _text(source.get("title", "集計表"))
        for headers, rows, offset in _panels(source, max_columns=4, cell_chars=64):
            # Tables use readable cell sizes instead of squeezing an arbitrarily
            # large grid onto one slide. Complete labels are on a preceding slide.
            prose(name + " / 列の説明", "\n".join(f"列 {offset + index}: {value}" for index, value in enumerate(headers[1:], 1)))
            for start in range(0, max(1, len(rows)), 3):
                item = slide(name + f" / 行 {start + 1} 以降")
                subset = rows[start:start + 3]
                table = item.shapes.add_table(len(subset) + 1, len(headers), Inches(.55), Inches(1.55), Inches(12.2), Inches(4.8)).table
                table.columns[0].width = Inches(.9)
                remaining = (12.2 - .9) / (len(headers) - 1)
                for column in list(table.columns)[1:]:
                    column.width = Inches(remaining)
                values = [["行", *[f"列 {offset + index}" for index in range(1, len(headers))]], *subset]
                for ri, row in enumerate(values):
                    for ci, value in enumerate(row):
                        cell = table.cell(ri, ci)
                        cell.text = _xml(value)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string("DCEFEB" if ri == 0 else "FFFFFF")
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name = OFFICE_FONT
                            paragraph.font.size = Pt(13)
                            paragraph.font.color.rgb = RGBColor.from_string("183147")
    for figure in bundle.get("figures", []):
        raw, w, h = _figure(figure)
        item = slide(_text(figure.get("title", "分析図")))
        scale = min(12.1 / w, 5.5 / h)
        width, height = w * scale, h * scale
        item.shapes.add_picture(BytesIO(raw), Inches((13.333 - width) / 2), Inches(1.35 + (5.5 - height) / 2), Inches(width), Inches(height))
        # Captions have their own pages so no explanatory text is clipped beneath
        # a landscape, time-layer plot or tall network figure.
        if figure.get("caption") or len(_text(figure.get("title", ""))) > 68:
            prose(_text(figure.get("title", "分析図")) + " / 図の説明", _text(figure.get("caption", "")))
    target = BytesIO()
    presentation.save(target)
    return target.getvalue()
