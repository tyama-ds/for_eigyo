import base64
from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from chat_app.attachments import MAX_FILE_BYTES, MAX_TEXT_CHARS, extract_attachment


def _saved(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


class AttachmentTests(unittest.TestCase):
    def test_text_encodings_and_filename(self):
        for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp932"):
            with self.subTest(encoding=encoding):
                result = extract_attachment("C:\\fakepath\\議事録.txt", "日本語の会議\r\n次の行".encode(encoding))
                self.assertEqual(result["name"], "議事録.txt")
                self.assertEqual(result["text"], "日本語の会議\n次の行")
                self.assertEqual(result["kind"], "text")

    def test_html_excludes_active_content(self):
        result = extract_attachment("page.html", b"<style>hidden</style><p>Hello &amp; world</p><script>alert(1)</script><p>Next</p>")
        self.assertIn("Hello & world", result["text"])
        self.assertIn("Next", result["text"])
        self.assertNotIn("hidden", result["text"])
        self.assertNotIn("alert", result["text"])

    def test_input_and_output_bounds(self):
        with self.assertRaisesRegex(ValueError, "20 MB"):
            extract_attachment("large.txt", b"a" * (MAX_FILE_BYTES + 1))
        result = extract_attachment("long.md", b"a" * (MAX_TEXT_CHARS + 1))
        self.assertEqual(len(result["text"]), MAX_TEXT_CHARS)
        self.assertIn("省略", result["warning"])

    def test_unsupported_empty_binary_and_invalid_encoding(self):
        for name, data in (("old.doc", b"123"), ("recording.mp3", b"123"), ("blank.txt", b""),
                           ("binary.txt", b"a\x00b"), ("invalid.txt", b"\x81")):
            with self.subTest(name=name), self.assertRaises(ValueError):
                extract_attachment(name, data)

    def test_docx_preserves_paragraph_and_table_order(self):
        from docx import Document

        document = Document()
        document.add_paragraph("先頭")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "項目"
        table.cell(0, 1).text = "金額"
        document.add_paragraph("末尾")
        result = extract_attachment("test.docx", _saved(document))
        self.assertEqual(result["text"], "先頭\n項目\t金額\n末尾")

    def test_xlsx_retains_formula_and_bounds_rows(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "集計"
        sheet.append(["品目", "数値", "合計"])
        sheet.append(["テスト", 10, "=SUM(B2:B2)"])
        sheet.cell(1001, 1, "範囲外")
        result = extract_attachment("test.xlsx", _saved(workbook))
        self.assertIn("[シート: 集計]", result["text"])
        self.assertIn("=SUM(B2:B2)", result["text"])
        self.assertNotIn("範囲外", result["text"])
        self.assertIn("1000行", result["warning"])

    def test_pptx_reads_text_tables_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "提案資料"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        slide.notes_slide.notes_text_frame.text = "発表者ノート"
        result = extract_attachment("test.pptx", _saved(presentation))
        self.assertIn("提案資料", result["text"])
        self.assertIn("A\tB", result["text"])
        self.assertIn("発表者ノート", result["text"])

    def test_xlsx_large_cells_respect_output_bound(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["a" * 32767] * 4)
        result = extract_attachment("large-cells.xlsx", _saved(workbook))
        self.assertEqual(len(result["text"]), MAX_TEXT_CHARS)
        self.assertIn("省略", result["warning"])

    def test_library_errors_do_not_expose_technical_details(self):
        with patch("chat_app.attachments._extract_pdf", side_effect=ValueError("internal secret path")):
            with self.assertRaisesRegex(ValueError, "読み込めません") as caught:
                extract_attachment("bad.pdf", b"%PDF")
            self.assertNotIn("internal", str(caught.exception))

    def test_pdf_blank_and_encrypted_rejected_clearly(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "OCR"):
            extract_attachment("scan.pdf", stream.getvalue())
        writer.encrypt("secret")
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "パスワード"):
            extract_attachment("protected.pdf", stream.getvalue())

    def test_pdf_extracts_page_text(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        stream = BytesIO()
        writer.write(stream)
        result = extract_attachment("hello.pdf", stream.getvalue())
        self.assertIn("Hello PDF", result["text"])
        self.assertIn("ページ 1", result["text"])

    def test_ooxml_zip_preflight(self):
        def archive_with(name, body=b"abc"):
            stream = BytesIO()
            with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr(name, body)
            return stream.getvalue()

        with self.assertRaisesRegex(ValueError, "ファイル名"):
            extract_attachment("evil.docx", archive_with("../outside.xml"))
        with patch("chat_app.attachments.MAX_ZIP_BYTES", 10):
            with self.assertRaisesRegex(ValueError, "展開後"):
                extract_attachment("bomb.xlsx", archive_with("large.xml", b"a" * 100))
        with patch("chat_app.attachments.MAX_ZIP_ENTRIES", 1):
            with self.assertRaisesRegex(ValueError, "ファイル数"):
                extract_attachment("entries.pptx", archive_with("a.xml"))
        with self.assertRaisesRegex(ValueError, "Office文書"):
            extract_attachment("corrupt.docx", b"not a zip")

    def test_image_resizes_and_validates_data_url(self):
        from PIL import Image

        original = Image.new("RGB", (3000, 1500), "navy")
        result = extract_attachment("photo.png", _saved_image(original, "PNG"))
        self.assertEqual(result["kind"], "image")
        encoded = base64.b64decode(result["data_url"].split(",", 1)[1])
        with Image.open(BytesIO(encoded)) as image:
            self.assertEqual(image.size, (2048, 1024))
            self.assertEqual(image.mode, "RGB")
        self.assertIn("縮小", result["warning"])
        with self.assertRaises(ValueError):
            extract_attachment("broken.png", b"not an image")

    def test_transparent_image_preserved(self):
        from PIL import Image

        result = extract_attachment("alpha.png", _saved_image(Image.new("RGBA", (10, 10), (0, 0, 0, 0)), "PNG"))
        self.assertEqual(result["mime"], "image/png")
        with Image.open(BytesIO(base64.b64decode(result["data_url"].split(",", 1)[1]))) as image:
            self.assertEqual(image.mode, "RGBA")

    def test_image_pixel_limit_and_format_mismatch(self):
        from PIL import Image

        picture = Image.new("RGB", (20, 20))
        with patch("chat_app.attachments.MAX_IMAGE_PIXELS", 100):
            with self.assertRaisesRegex(ValueError, "画素数"):
                extract_attachment("too-large.png", _saved_image(picture, "PNG"))
        with self.assertRaisesRegex(ValueError, "PNG"):
            extract_attachment("renamed.png", _saved_image(picture, "BMP"))


def _saved_image(image, format_name):
    stream = BytesIO()
    image.save(stream, format=format_name)
    return stream.getvalue()


if __name__ == "__main__":
    unittest.main()
