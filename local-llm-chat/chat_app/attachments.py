"""Extract user attachments in memory, with bounded document and image output."""

from __future__ import annotations

import base64
from html.parser import HTMLParser
from io import BytesIO
from pathlib import PurePosixPath
import re
import warnings
from zipfile import BadZipFile, ZipFile


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 80_000
MAX_ZIP_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 5_000
MAX_IMAGE_PIXELS = 40_000_000
MAX_IMAGE_EDGE = 2_048
MAX_SHEETS = 20
MAX_SHEET_ROWS = 1_000
MAX_SHEET_COLUMNS = 50
MAX_PDF_PAGES = 500
TEXT_EXTENSIONS = {
    ".txt", ".text", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
    ".jsonl", ".ndjson", ".xml", ".html", ".htm", ".log", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".properties", ".env", ".sql", ".py",
    ".js", ".jsx", ".ts", ".tsx", ".css", ".scss", ".sass", ".less", ".vue",
    ".svelte", ".java", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".go",
    ".rs", ".rb", ".php", ".swift", ".kt", ".kts", ".sh", ".bash", ".zsh",
    ".ps1", ".bat", ".cmd", ".r", ".lua", ".tex", ".bib", ".ipynb",
    ".svg", ".srt", ".vtt", ".dockerfile", ".gitignore", ".editorconfig",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}


class _AttachmentError(ValueError):
    """An expected, already localized validation failure."""


class _TextBuffer:
    def __init__(self) -> None:
        self.parts: list[str] = []
        self.length = 0
        self.truncated = False

    def add(self, text: str) -> bool:
        if not text:
            return self.length < MAX_TEXT_CHARS
        remaining = MAX_TEXT_CHARS - self.length
        if self.parts and remaining > 0:
            self.parts.append("\n")
            self.length += 1
            remaining -= 1
        if len(text) > remaining:
            self.truncated = True
        self.parts.append(text[:remaining])
        self.length += min(len(text), remaining)
        return self.length < MAX_TEXT_CHARS

    def value(self) -> str:
        return "".join(self.parts)


class _HTMLText(HTMLParser):
    BLOCKS = {"p", "div", "br", "li", "tr", "section", "article", "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in {"script", "style"}:
            self.hidden += 1
        elif not self.hidden and tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden and tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def _safe_name(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return (name or "attachment")[:240]


def _preflight_ooxml(data: bytes) -> None:
    """Inspect archive metadata; never extract archive entries to disk."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise _AttachmentError("文書内のファイル数が上限（5,000件）を超えています。")
            total = 0
            names = set()
            for entry in entries:
                total += entry.file_size
                if total > MAX_ZIP_BYTES:
                    raise _AttachmentError("文書の展開後サイズが上限（100 MB）を超えています。")
                name = entry.filename.replace("\\", "/")
                path = PurePosixPath(name)
                if name.startswith("/") or ".." in path.parts or ":" in name or "\x00" in name:
                    raise _AttachmentError("文書内に安全に読み込めないファイル名があります。")
                if entry.flag_bits & 1:
                    raise _AttachmentError("パスワード保護された文書には対応していません。")
                if name in names:
                    raise _AttachmentError("文書内に重複するファイル名があります。")
                names.add(name)
            if "[Content_Types].xml" not in names:
                raise _AttachmentError("Office文書の形式が正しくありません。")
    except BadZipFile as exc:
        raise _AttachmentError("Office文書が破損しているか、パスワードで保護されています。") from exc


def _decode_text(data: bytes) -> str:
    encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp932")
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeError:
            continue
        controls = sum(ord(char) < 32 and char not in "\t\n\r\f" for char in text)
        if "\x00" in text or controls > max(2, len(text) // 100):
            raise _AttachmentError("バイナリデータはテキストとして読み込めません。")
        return text.replace("\r\n", "\n").replace("\r", "\n")
    raise _AttachmentError("文字コードを判別できません。UTF-8、UTF-16（BOM付き）、またはCP932で保存してください。")


def _extract_pdf(data: bytes, output: _TextBuffer, notices: list[str]) -> None:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise _AttachmentError("パスワード保護されたPDFには対応していません。")
    count = len(reader.pages)
    if count > MAX_PDF_PAGES:
        notices.append(f"PDFは先頭{MAX_PDF_PAGES}ページまで読み込みました。")
    found = False
    empty_pages = 0
    for index in range(min(count, MAX_PDF_PAGES)):
        content = (reader.pages[index].extract_text() or "").strip()
        if content:
            found = True
            if not output.add(f"[ページ {index + 1}]\n{content}"):
                if index + 1 < count:
                    output.truncated = True
                break
        else:
            empty_pages += 1
    if not found:
        raise _AttachmentError("PDFからテキストを抽出できませんでした。画像のみのPDF・スキャンPDFのOCRには対応していません。")
    if empty_pages:
        notices.append(f"{empty_pages}ページには抽出できるテキストがありませんでした。画像部分のOCRは行いません。")


def _extract_docx(data: bytes, output: _TextBuffer, notices: list[str]) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(data))
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            text = Paragraph(element, document).text
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            for row in table.rows:
                if not output.add("\t".join(cell.text for cell in row.cells)):
                    output.truncated = True
                    return
            continue
        else:
            continue
        if not output.add(text):
            output.truncated = True
            return
    seen = set()
    for section in document.sections:
        for region in (section.header, section.footer):
            if region.part.partname in seen:
                continue
            seen.add(region.part.partname)
            for paragraph in region.paragraphs:
                if not output.add(paragraph.text):
                    output.truncated = True
                    return


def _extract_xlsx(data: bytes, output: _TextBuffer, notices: list[str]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            notices.append(f"Excelは先頭{MAX_SHEETS}シートまで読み込みました。")
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            if not output.add(f"[シート: {sheet.title}]"):
                output.truncated = True
                return
            if (sheet.max_row or 0) > MAX_SHEET_ROWS or (sheet.max_column or 0) > MAX_SHEET_COLUMNS:
                notices.append(f"「{sheet.title}」は先頭{MAX_SHEET_ROWS}行・{MAX_SHEET_COLUMNS}列まで読み込みました。")
            for row in sheet.iter_rows(max_row=min(sheet.max_row or MAX_SHEET_ROWS, MAX_SHEET_ROWS),
                                       max_col=min(sheet.max_column or MAX_SHEET_COLUMNS, MAX_SHEET_COLUMNS),
                                       values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells and not output.add("\t".join(cells)):
                    output.truncated = True
                    return
        notices.append("Excelの数式は計算せず、数式そのものを読み込んでいます。")
    finally:
        workbook.close()


def _extract_pptx(data: bytes, output: _TextBuffer, notices: list[str]) -> None:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))

    def add_shapes(shapes, depth: int = 0) -> bool:
        if depth > 20:
            notices.append("深く入れ子になった図形の一部を省略しました。")
            return True
        for shape in shapes:
            if shape.has_text_frame and not output.add(shape.text):
                return False
            if shape.has_table:
                for row in shape.table.rows:
                    if not output.add("\t".join(cell.text for cell in row.cells)):
                        return False
            if hasattr(shape, "shapes") and not add_shapes(shape.shapes, depth + 1):
                return False
        return True

    for index, slide in enumerate(presentation.slides, 1):
        if not output.add(f"[スライド {index}]") or not add_shapes(slide.shapes):
            output.truncated = True
            return
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip() and not output.add(f"[ノート]\n{notes.text}"):
                output.truncated = True
                return


def _extract_image(name: str, data: bytes) -> dict:
    from PIL import Image, ImageOps

    notices = []
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as probe:
            if probe.format not in {"PNG", "JPEG", "WEBP", "GIF"}:
                raise _AttachmentError("画像はPNG・JPEG・WebP・GIFに対応しています。")
            if probe.width * probe.height > MAX_IMAGE_PIXELS:
                raise _AttachmentError("画像の画素数が上限（4,000万画素）を超えています。")
            probe.verify()
        with Image.open(BytesIO(data)) as source:
            if getattr(source, "n_frames", 1) > 1:
                notices.append("アニメーション画像は先頭フレームを使用します。")
            image = ImageOps.exif_transpose(source)
            if max(image.size) > MAX_IMAGE_EDGE:
                image.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE), Image.Resampling.LANCZOS)
                notices.append("画像の長辺を2,048ピクセル以内に縮小しました。")
            if image.mode in {"RGBA", "LA"} or "transparency" in image.info:
                image = image.convert("RGBA")
                mime, format_name = "image/png", "PNG"
            else:
                image = image.convert("RGB")
                mime, format_name = "image/jpeg", "JPEG"
            encoded = BytesIO()
            image.save(encoded, format=format_name, **({"quality": 88} if format_name == "JPEG" else {}))
    result = {"name": name, "kind": "image", "size": len(data), "text": "", "mime": mime,
              "data_url": f"data:{mime};base64,{base64.b64encode(encoded.getvalue()).decode('ascii')}"}
    notices.append("画像の理解には画像入力対応モデルが必要です。OCRによるテキスト抽出は行いません。")
    result["warning"] = " ".join(notices)
    return result


def extract_attachment(filename: str, data: bytes) -> dict:
    """Return text or a normalized image for a chat message, without disk writes.

    All expected input failures are exposed as user-safe Japanese ValueErrors.
    The size field always describes the original input, not extracted output.
    """
    name = _safe_name(filename)
    if len(data) > MAX_FILE_BYTES:
        raise _AttachmentError("添付ファイルは1件あたり20 MB以内にしてください。")
    if not data:
        raise _AttachmentError("空のファイルは読み込めません。")
    extension = PurePosixPath(name).suffix.lower()
    if not extension and name.lower() in {"dockerfile", "makefile", "license", "readme", ".env", ".gitignore", ".editorconfig"}:
        extension = ".txt"
    if extension not in TEXT_EXTENSIONS | IMAGE_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf"}:
        raise _AttachmentError("このファイル形式には対応していません。テキスト、PDF、DOCX、XLSX、PPTX、PNG、JPEG、WebP、GIFをご利用ください。旧Office形式・音声・動画・ZIPには対応していません。")
    output = _TextBuffer()
    notices: list[str] = []
    try:
        if extension in IMAGE_EXTENSIONS:
            return _extract_image(name, data)
        if extension in OFFICE_EXTENSIONS:
            _preflight_ooxml(data)
        if extension in TEXT_EXTENSIONS:
            text = _decode_text(data)
            if extension in {".html", ".htm"}:
                parser = _HTMLText()
                parser.feed(text)
                parser.close()
                text = re.sub(r"\n[ \t]*\n+", "\n\n", "".join(parser.parts)).strip()
            output.add(text)
        elif extension == ".pdf":
            _extract_pdf(data, output, notices)
        elif extension == ".docx":
            _extract_docx(data, output, notices)
        elif extension == ".xlsx":
            _extract_xlsx(data, output, notices)
        elif extension == ".pptx":
            _extract_pptx(data, output, notices)
    except _AttachmentError:
        raise
    except Exception as exc:
        raise _AttachmentError("ファイルを読み込めませんでした。破損、形式の不一致、パスワード保護の有無を確認してください。") from exc
    text = output.value()
    if not text.strip():
        raise _AttachmentError("ファイルから読み込めるテキストが見つかりませんでした。画像や図表のOCRには対応していません。")
    if output.truncated:
        notices.append(f"抽出テキストは先頭{MAX_TEXT_CHARS:,}文字で省略しました。")
    result = {"name": name, "kind": "text", "size": len(data), "text": text}
    if notices:
        result["warning"] = " ".join(dict.fromkeys(notices))
    return result
