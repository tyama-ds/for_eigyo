"""CSV / XLSX の読み込みと列型推論。

- 標準ライブラリのみで動く（openpyxl が入っていれば XLSX 読み込みに優先利用）
- CSV は utf-8 / utf-8-sig / cp932(Shift_JIS) / euc_jp を自動判別
- XLSX は zip + XML を直接パースする最小リーダーを内蔵（共有文字列・日付書式に対応）
"""
from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime, timedelta
from xml.etree import ElementTree as ET

MISSING_TOKENS = {"", "nan", "NaN", "NAN", "na", "NA", "n/a", "N/A", "null", "NULL",
                  "None", "none", "-", "—", "－", "#N/A", "#VALUE!", "#DIV/0!", "#REF!"}
MAX_PREVIEW_ROWS = 50
HIST_BINS = 20

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_EXCEL_EPOCH = datetime(1899, 12, 30)
_DATE_NUMFMT_IDS = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | set(range(50, 59))
_DATE_RE = re.compile(r"^\d{4}[-/年.]\d{1,2}[-/月.]\d{1,2}(日)?([ T]\d{1,2}:\d{2}(:\d{2})?)?$")
_CJK_RE = re.compile(r"[぀-ヿ㐀-鿿]")


class DataError(Exception):
    """読み込みや解釈の失敗（ユーザーに表示するメッセージ）。"""


# ---------------------------------------------------------------- 値の解釈

def is_missing(value) -> bool:
    if value is None:
        return True
    return str(value).strip() in MISSING_TOKENS


def parse_number(value) -> float | None:
    """'1,234' '¥1,200' '12%' '１２３' などを数値にする。数値でなければ None。"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    s = unicodedata.normalize("NFKC", str(value)).strip()
    if s in MISSING_TOKENS:
        return None
    s = s.replace(",", "").replace(" ", "")
    s = s.lstrip("¥$€£").rstrip("%")
    if s.startswith("(") and s.endswith(")"):        # 会計表記の負数 (1,200)
        s = "-" + s[1:-1]
    if not s or s in {"+", "-", "."}:
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    if math.isinf(v) or math.isnan(v):
        return None
    return v


def looks_like_date(value: str) -> bool:
    return bool(_DATE_RE.match(value.strip()))


# ---------------------------------------------------------------- CSV

def _decode_bytes(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "cp932", "euc_jp"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8(置換あり)"


def _sniff_delimiter(text: str, filename: str) -> str:
    if filename.lower().endswith((".tsv", ".tab")):
        return "\t"
    head = "\n".join(text.splitlines()[:30])
    try:
        return csv.Sniffer().sniff(head, delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def read_csv_bytes(data: bytes, filename: str = "data.csv") -> dict:
    text, enc = _decode_bytes(data)
    delim = _sniff_delimiter(text, filename)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [row for row in reader]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise DataError("CSV にデータ行がありません")
    header, body = rows[0], rows[1:]
    return _make_table(header, body, filename, meta={"encoding": enc, "delimiter": delim})


# ---------------------------------------------------------------- XLSX（最小リーダー）

def _col_index(ref: str) -> int:
    n = 0
    for ch in ref:
        if ch.isalpha():
            n = n * 26 + (ord(ch.upper()) - 64)
        else:
            break
    return max(n - 1, 0)


def _is_date_format(code: str) -> bool:
    code = re.sub(r'"[^"]*"|\[[^\]]*\]', "", code)   # 引用・[$-ja-JP] 等を除く
    return bool(re.search(r"[ymdhs]", code.lower())) and not re.search(r"[0#?]", code)


def _excel_serial_to_str(v: float) -> str:
    try:
        dt = _EXCEL_EPOCH + timedelta(days=v)
    except (OverflowError, ValueError):
        return str(v)
    if abs(v - round(v)) < 1e-9:
        return dt.strftime("%Y-%m-%d")
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _xlsx_sheet_list(zf: zipfile.ZipFile) -> list[tuple[str, str]]:
    """[(シート名, zip 内パス)] を workbook.xml と rels から得る。"""
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    target_of = {}
    for rel in rels.iter(rel_ns + "Relationship"):
        target = rel.get("Target", "")
        if target.startswith("/"):
            target = target[1:]
        elif not target.startswith("xl/"):
            target = "xl/" + target
        target_of[rel.get("Id")] = target
    doc_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    sheets = []
    for sh in wb.iter(_NS + "sheet"):
        rid = sh.get(doc_ns + "id")
        path = target_of.get(rid)
        if path and path in zf.namelist():
            sheets.append((sh.get("name", "Sheet"), path))
    if not sheets:
        raise DataError("XLSX にワークシートが見つかりません")
    return sheets


def _xlsx_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    out = []
    for _ev, el in ET.iterparse(io.BytesIO(zf.read("xl/sharedStrings.xml"))):
        if el.tag == _NS + "si":
            out.append("".join(t.text or "" for t in el.iter(_NS + "t")))
            el.clear()
    return out


def _xlsx_date_styles(zf: zipfile.ZipFile) -> set[int]:
    """日付書式を持つ cellXfs のインデックス集合。"""
    if "xl/styles.xml" not in zf.namelist():
        return set()
    root = ET.fromstring(zf.read("xl/styles.xml"))
    custom = {}
    for nf in root.iter(_NS + "numFmt"):
        try:
            custom[int(nf.get("numFmtId"))] = nf.get("formatCode", "")
        except (TypeError, ValueError):
            pass
    date_styles = set()
    xfs = root.find(_NS + "cellXfs")
    if xfs is None:
        return date_styles
    for i, xf in enumerate(xfs.findall(_NS + "xf")):
        try:
            fid = int(xf.get("numFmtId", "0"))
        except ValueError:
            continue
        if fid in _DATE_NUMFMT_IDS or (fid in custom and _is_date_format(custom[fid])):
            date_styles.add(i)
    return date_styles


def _xlsx_read_sheet(zf: zipfile.ZipFile, path: str, shared: list[str],
                     date_styles: set[int]) -> list[list[str]]:
    rows: list[list[str]] = []
    for _ev, el in ET.iterparse(io.BytesIO(zf.read(path))):
        if el.tag != _NS + "row":
            continue
        cells: dict[int, str] = {}
        for c in el.findall(_NS + "c"):
            ref = c.get("r", "")
            idx = _col_index(ref) if ref else len(cells)
            t = c.get("t", "n")
            v_el = c.find(_NS + "v")
            if t == "inlineStr":
                val = "".join(x.text or "" for x in c.iter(_NS + "t"))
            elif v_el is None or v_el.text is None:
                val = ""
            elif t == "s":
                try:
                    val = shared[int(v_el.text)]
                except (ValueError, IndexError):
                    val = ""
            elif t == "b":
                val = "TRUE" if v_el.text.strip() == "1" else "FALSE"
            elif t in ("str", "e"):
                val = v_el.text
            else:
                val = v_el.text.strip()
                style = c.get("s")
                if style is not None and int(style) in date_styles:
                    num = parse_number(val)
                    if num is not None:
                        val = _excel_serial_to_str(num)
            cells[idx] = val
        if cells:
            width = max(cells) + 1
            rows.append([cells.get(i, "") for i in range(width)])
        el.clear()
    return rows


def read_xlsx_bytes(data: bytes, filename: str = "data.xlsx", sheet: str | None = None) -> dict:
    try:
        return _read_xlsx_openpyxl(data, filename, sheet)
    except ImportError:
        pass
    except DataError:
        raise
    except Exception:  # noqa: BLE001 — openpyxl が読めないファイルは内蔵リーダーで再試行
        pass
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise DataError("XLSX として開けません（.xls 旧形式は未対応。xlsx で保存し直してください）") from e
    with zf:
        sheets = _xlsx_sheet_list(zf)
        names = [n for n, _ in sheets]
        chosen = sheets[0]
        if sheet:
            for n, p in sheets:
                if n == sheet:
                    chosen = (n, p)
                    break
            else:
                raise DataError(f"シート '{sheet}' が見つかりません")
        shared = _xlsx_shared_strings(zf)
        date_styles = _xlsx_date_styles(zf)
        rows = _xlsx_read_sheet(zf, chosen[1], shared, date_styles)
    return _rows_to_table(rows, filename, meta={"sheet": chosen[0], "sheets": names, "reader": "builtin"})


def _read_xlsx_openpyxl(data: bytes, filename: str, sheet: str | None) -> dict:
    import openpyxl  # noqa: F401  (ImportError は呼び出し側で捕捉)

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = wb.sheetnames
    if sheet and sheet not in names:
        raise DataError(f"シート '{sheet}' が見つかりません")
    ws = wb[sheet] if sheet else wb[names[0]]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([_cell_to_str(v) for v in row])
    wb.close()
    return _rows_to_table(rows, filename, meta={"sheet": ws.title, "sheets": names, "reader": "openpyxl"})


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ---------------------------------------------------------------- 共通

def _rows_to_table(rows: list[list[str]], filename: str, meta: dict) -> dict:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise DataError("データ行がありません")
    return _make_table(rows[0], rows[1:], filename, meta)


def _make_table(header: list, body: list[list], filename: str, meta: dict) -> dict:
    width = max([len(header)] + [len(r) for r in body])
    columns = []
    seen: Counter = Counter()
    for i in range(width):
        name = str(header[i]).strip() if i < len(header) else ""
        if not name:
            name = f"列{i + 1}"
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"
        columns.append(name)
    norm_rows = []
    for r in body:
        r = [("" if c is None else str(c)) for c in r]
        if len(r) < width:
            r = r + [""] * (width - len(r))
        norm_rows.append(r[:width])
    if not norm_rows:
        raise DataError("ヘッダー行しかありません（データ行が必要です）")
    table = {"name": filename, "columns": columns, "rows": norm_rows, "n_rows": len(norm_rows)}
    table.update(meta)
    return table


def read_table_bytes(data: bytes, filename: str, sheet: str | None = None) -> dict:
    """拡張子で CSV / XLSX を判定して読み込む。"""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xltx")):
        return read_xlsx_bytes(data, filename, sheet)
    if lower.endswith(".xls"):
        raise DataError(".xls（旧 Excel 形式）は未対応です。.xlsx で保存し直してください")
    if data[:2] == b"PK":   # 拡張子が違っても zip なら xlsx 扱い
        return read_xlsx_bytes(data, filename, sheet)
    return read_csv_bytes(data, filename)


# ---------------------------------------------------------------- 列型推論・統計

def histogram(values: list[float], bins: int = HIST_BINS) -> dict:
    if not values:
        return {"edges": [], "counts": []}
    lo, hi = min(values), max(values)
    if lo == hi:
        return {"edges": [lo, hi], "counts": [len(values)]}
    step = (hi - lo) / bins
    counts = [0] * bins
    for v in values:
        i = int((v - lo) / step)
        if i >= bins:
            i = bins - 1
        counts[i] += 1
    edges = [lo + step * i for i in range(bins + 1)]
    return {"edges": edges, "counts": counts}


def _quantile(sorted_vals: list[float], q: float) -> float:
    if not sorted_vals:
        return float("nan")
    pos = (len(sorted_vals) - 1) * q
    lo = int(math.floor(pos))
    hi = min(lo + 1, len(sorted_vals) - 1)
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (pos - lo)


def analyze_column(name: str, values: list[str], index: int = 0) -> dict:
    n = len(values)
    present = [v for v in values if not is_missing(v)]
    missing = n - len(present)
    info: dict = {"name": name, "index": index, "n": n, "missing": missing,
                  "missing_pct": round(100.0 * missing / n, 1) if n else 0.0}
    if not present:
        info.update({"type": "empty", "unique": 0, "stats": {}, "samples": []})
        return info
    counter = Counter(v.strip() for v in present)
    unique = len(counter)
    info["unique"] = unique
    info["samples"] = [k for k, _ in counter.most_common(5)]

    numbers = [parse_number(v) for v in present]
    num_ok = [x for x in numbers if x is not None]
    num_ratio = len(num_ok) / len(present)
    date_ratio = sum(1 for v in present if looks_like_date(v)) / len(present)
    lengths = [len(v.strip()) for v in present]
    avg_len = sum(lengths) / len(lengths)
    cjk_ratio = sum(1 for v in present if _CJK_RE.search(v)) / len(present)
    space_ratio = sum(1 for v in present if " " in v.strip() or "　" in v.strip()) / len(present)

    if num_ratio >= 0.9 and date_ratio < 0.5:
        ctype = "numeric"
    elif date_ratio >= 0.9:
        ctype = "datetime"
    elif (unique <= 20 or unique <= 0.05 * len(present)) and avg_len < 30:
        ctype = "categorical"
    elif avg_len >= 12 or space_ratio >= 0.3 or (cjk_ratio >= 0.5 and avg_len >= 6):
        ctype = "text"
    elif unique == len(present) and unique > 20:
        ctype = "id"
    else:
        ctype = "categorical"
    info["type"] = ctype

    stats: dict = {}
    if ctype == "numeric":
        s = sorted(num_ok)
        mean = sum(s) / len(s)
        var = sum((x - mean) ** 2 for x in s) / max(len(s) - 1, 1)
        stats = {"min": s[0], "max": s[-1], "mean": mean, "std": math.sqrt(var),
                 "median": _quantile(s, 0.5), "q1": _quantile(s, 0.25), "q3": _quantile(s, 0.75),
                 "hist": histogram(s), "non_numeric": len(present) - len(num_ok),
                 "integer_like": all(abs(x - round(x)) < 1e-9 for x in s)}
    elif ctype in ("categorical", "id", "datetime"):
        stats = {"top": [{"value": k, "count": c} for k, c in counter.most_common(12)]}
        if ctype == "datetime":
            stats["min"] = min(present)
            stats["max"] = max(present)
    elif ctype == "text":
        s = sorted(lengths)
        stats = {"avg_len": avg_len, "max_len": s[-1], "min_len": s[0],
                 "p95_len": _quantile(s, 0.95), "len_hist": histogram([float(x) for x in s], 15),
                 "cjk_ratio": cjk_ratio}
    info["stats"] = stats
    return info


def analyze_table(table: dict) -> list[dict]:
    cols = table["columns"]
    rows = table["rows"]
    out = []
    for i, name in enumerate(cols):
        out.append(analyze_column(name, [r[i] for r in rows], i))
    return out


def suggest_task(col: dict) -> str:
    """目的変数候補の列情報からタスク種別を推定する。"""
    if col["type"] == "numeric":
        n = max(col["n"] - col["missing"], 1)
        if col["unique"] > 20 or col["unique"] / n > 0.2:
            return "regression"
        return "classification"
    return "classification"


def suggest_roles(analysis: list[dict], target: str | None = None) -> dict:
    """列ごとの既定ロール（target / text / numeric / categorical / ignore）。"""
    if not analysis:
        return {"target": None, "task": "classification", "roles": {}}
    if target is None or target not in {c["name"] for c in analysis}:
        # 末尾から最初に見つかった「使えそうな」列を目的変数候補にする
        target = analysis[-1]["name"]
        for c in reversed(analysis):
            if c["type"] in ("numeric", "categorical") and c["unique"] >= 2:
                target = c["name"]
                break
    roles = {}
    tcol = None
    for c in analysis:
        if c["name"] == target:
            roles[c["name"]] = "target"
            tcol = c
        elif c["type"] in ("numeric", "categorical", "text"):
            roles[c["name"]] = c["type"]
        else:
            roles[c["name"]] = "ignore"
    task = suggest_task(tcol) if tcol else "classification"
    return {"target": target, "task": task, "roles": roles}


def table_summary(table: dict, analysis: list[dict] | None = None) -> dict:
    analysis = analysis or analyze_table(table)
    rows = table["rows"]
    total_cells = len(rows) * len(table["columns"])
    missing_cells = sum(c["missing"] for c in analysis)
    return {
        "name": table.get("name"),
        "n_rows": len(rows),
        "n_cols": len(table["columns"]),
        "columns": table["columns"],
        "sheet": table.get("sheet"),
        "sheets": table.get("sheets"),
        "encoding": table.get("encoding"),
        "reader": table.get("reader"),
        "missing_pct": round(100.0 * missing_cells / total_cells, 2) if total_cells else 0.0,
        "analysis": analysis,
        "preview": rows[:MAX_PREVIEW_ROWS],
        "suggest": suggest_roles(analysis),
    }
